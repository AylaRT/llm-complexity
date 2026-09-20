"""Read the gold standard and the parser baselines, and score counts against
them.

Everything comes from the "All" sheet of results_new.xlsx: column B holds the
text id, then the same five units in the same order three times over, for the
MSD manual counts (C-G), TAASSC (I-M) and L2SCA (O-S). Levels are the "level N"
headers in column A, each grouping the 20 texts below it.

Scoring is always per unit type, never a single blended score.
"""

import re
from pathlib import Path

import numpy as np
import openpyxl
from scipy import stats

from src.parse import UNIT_KEYS

_FIRST_DATA_ROW = 3
_ID_COLUMN = 2
_LEVEL_COLUMN = 1

# First column (1-based) of each source block; the five units follow in the
# order given by _BLOCK_UNITS.
_BLOCK_START = {"gold": 3, "taassc": 9, "l2sca": 15}
_BLOCK_UNITS = ["sentences", "t_units", "clauses", "dependent_clauses", "words"]


def _columns(source: str) -> dict[str, int]:
    start = _BLOCK_START[source]
    return {unit: start + i for i, unit in enumerate(_BLOCK_UNITS)}


def _number_in(value) -> int | None:
    """The first integer in a cell such as "text 12" or "level 3"."""
    if value is None:
        return None
    match = re.search(r"(\d+)", str(value))
    return int(match.group(1)) if match else None


def _rows(sheet):
    """Yield (text_id, row index) for every text row in the sheet."""
    for row in range(_FIRST_DATA_ROW, sheet.max_row + 1):
        label = sheet.cell(row, _ID_COLUMN).value
        if not label or "text" not in str(label).lower():
            continue
        text_id = _number_in(label)
        if text_id is not None:
            yield text_id, row


def _read_block(path: Path, source: str, cast) -> dict[int, dict]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook["All"]
    columns = _columns(source)
    counts = {
        text_id: {
            unit: cast(sheet.cell(row, col).value)
            for unit, col in columns.items()
        }
        for text_id, row in _rows(sheet)
    }
    workbook.close()
    return counts


def _as_int(value) -> int:
    return round(value) if value is not None else 0


def _as_float(value) -> float:
    return float(value) if value is not None else 0.0


def load_gold(path: Path) -> dict[int, dict]:
    """MSD manual counts as {text_id: {unit: count}}."""
    return _read_block(path, "gold", _as_int)


def load_baselines(path: Path) -> dict[str, dict[int, dict]]:
    """TAASSC and L2SCA counts as {source: {text_id: {unit: count}}}."""
    return {name: _read_block(path, name, _as_float)
            for name in ("taassc", "l2sca")}


def load_levels(path: Path) -> dict[int, int]:
    """Proficiency level 1-4 per text, from the level headers in column A."""
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook["All"]

    levels: dict[int, int] = {}
    current = None
    for row in range(_FIRST_DATA_ROW, sheet.max_row + 1):
        label = sheet.cell(row, _LEVEL_COLUMN).value
        if label and "level" in str(label).lower():
            level = _number_in(label)
            if level is not None:
                current = level
        text_label = sheet.cell(row, _ID_COLUMN).value
        if text_label and "text" in str(text_label).lower() and current is not None:
            text_id = _number_in(text_label)
            if text_id is not None:
                levels[text_id] = current

    workbook.close()
    return levels


def score_unit(
    predicted: dict[int, float],
    reference: dict[int, float],
) -> dict:
    """Compare one unit type over the texts present in both.

    n records how many texts were scored: a condition that failed on hard
    texts is scored on an easier subset, so report it with the metrics.
    """
    common = sorted(set(predicted) & set(reference))
    if len(common) < 3:
        return {"n": len(common), "error": "too few texts to score"}

    pred = np.array([predicted[tid] for tid in common])
    ref = np.array([reference[tid] for tid in common])
    errors = pred - ref
    abs_errors = np.abs(errors)

    r, p = stats.pearsonr(pred, ref)

    return {
        "n": len(common),
        "pearson_r": round(float(r), 4),
        "pearson_p": float(p),
        "mean_error": round(float(abs_errors.mean()), 2),
        "mean_bias": round(float(errors.mean()), 2),
        "rmse": round(float(np.sqrt((errors ** 2).mean())), 2),
        "max_error": int(abs_errors.max()),
    }


def score_all(
    parsed: dict[int, dict[str, int]],
    gold: dict[int, dict],
) -> dict[str, dict]:
    """Score every unit type, keyed by unit."""
    results: dict[str, dict] = {}
    for unit in UNIT_KEYS:
        predicted = {tid: counts[unit] for tid, counts in parsed.items()
                     if unit in counts}
        reference = {tid: counts[unit] for tid, counts in gold.items()
                     if unit in counts}
        results[unit] = score_unit(predicted, reference)
    return results
