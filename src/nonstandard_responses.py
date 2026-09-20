"""Scan the saved responses for non-standard output.

This is the sweep behind the manual corrections in data/manual_overrides.json:
it collects the cases where a model answered in a shape the parser reads
literally but a reader would not, so each can be inspected and either corrected
or left alone. It changes nothing; it writes one xlsx row per pattern found.

Patterns, by format:

  counts  COUNTS_EXTRA_TEXT   anything beyond the five HEADER: N lines
  lists   DC_MIXED_NONE       real clauses alongside a none-equivalent line
          COMMENTARY          a line that reads as commentary, not an item
          INSTRUCTION_ECHO    a prompt instruction repeated as an item
  both    COUNT_OUTLIER_HIGH  a count far above the median across conditions
          COUNT_OUTLIER_LOW   a count far below it

Numbered and bulleted prefixes are not flagged: the parser strips those when
they are uniform across a block, and it does so for every response.
"""

import argparse
import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config.settings import CORPUS_PATH, OUTPUTS_DIR, PHASES, RESULTS_DIR
from src.overrides import load_overrides
from src.parse import HEADERS, UNIT_KEYS, parse, parse_dc_lines, split_blocks

_COMMENTARY_RE = re.compile(
    r"(^note:\s|^total\s*(words|sentences|t-units|clauses|count)|"
    r"^i\s+(count|found|identified)|^count:\s*\d|"
    r"^summary\s*:|^explanation\s*:|"
    r"^here are the|^the following are|^listed below|^below are the)",
    re.IGNORECASE,
)

_INSTRUCTION_ECHO_RE = re.compile(
    r"(if there are none,?\s*write|list each\b.*\bone per line|"
    r"comma[\s-]separated|separated by commas)",
    re.IGNORECASE,
)

_PHASE_SIZES = {"pilot": "small", "phase1": "small", "phase2": "large"}

# An outlier has to be far from the median in relative terms and at least
# _MIN_ABSOLUTE_GAP away from it, so that small units do not flag on noise.
# Where the median is itself small the ratio says little, and only a large
# absolute gap counts.
_MIN_CONDITIONS = 4
_MIN_ABSOLUTE_GAP = 10
_SMALL_MEDIAN = 5
_SMALL_MEDIAN_GAP = 15
_HIGH_RATIO = 5
_LOW_RATIO = 0.2

COLUMNS = [
    "phase", "model", "size", "reasoning", "format", "condition",
    "text_id", "rep", "attempt", "pattern",
    *UNIT_KEYS,
    *[f"override_{unit}" for unit in UNIT_KEYS],
    "output", "text",
]

_COLUMN_WIDTHS = {
    "phase": 8, "model": 16, "size": 6, "reasoning": 10, "format": 7,
    "condition": 32, "text_id": 7, "rep": 4, "attempt": 7, "pattern": 40,
    "words": 7, "sentences": 9, "t_units": 7, "clauses": 8,
    "dependent_clauses": 16, "output": 60, "text": 60,
}


def _load_texts(corpus_path: Path) -> dict[str, str]:
    with open(corpus_path, encoding="utf-8", newline="") as f:
        return {
            text_id: row["text"]
            for row in csv.DictReader(f, delimiter="\t")
            if (text_id := row.get("text_id") or row.get("pilot_id"))
        }


def _split_condition(condition: str) -> tuple[str, str]:
    """Split a condition name into (model, reasoning)."""
    if "_reasoning-" in condition:
        model, reasoning = condition.rsplit("_reasoning-", 1)
        return model, reasoning
    return condition, "unknown"


def _block_excerpt(unit: str, raw: str) -> str:
    """One block of a response, with its header, for the reviewer to read."""
    blocks = split_blocks(raw)
    if blocks is None:
        return raw
    return f"{HEADERS[unit]}\n{blocks.get(unit, '')}"


def _non_empty_lines(block: str) -> list[str]:
    return [ln.strip() for ln in block.splitlines() if ln.strip()]


def _detect_counts_patterns(raw: str) -> list[tuple[str, str]]:
    headers = tuple(HEADERS.values())
    extra = [ln for ln in _non_empty_lines(raw)
             if not ln.upper().startswith(headers)]
    return [("COUNTS_EXTRA_TEXT", raw)] if extra else []


def _detect_lists_patterns(raw: str) -> list[tuple[str, str]]:
    blocks = split_blocks(raw)
    if blocks is None:
        return []

    patterns = []

    dc_block = blocks.get("dependent_clauses", "")
    kept = parse_dc_lines(dc_block)
    if kept and len(_non_empty_lines(dc_block)) > len(kept):
        patterns.append(
            ("DC_MIXED_NONE", f"{HEADERS['dependent_clauses']}\n{dc_block}"))

    for unit in UNIT_KEYS:
        block = blocks.get(unit, "")
        excerpt = f"{HEADERS[unit]}\n{block}"
        if unit == "words":
            # The words block is one comma-separated line, so only its first
            # line can be read as commentary.
            first_line = block.split("\n")[0].strip() if block else ""
            if first_line and _COMMENTARY_RE.match(first_line):
                patterns.append(("COMMENTARY", excerpt))
            continue
        for line in _non_empty_lines(block):
            if _COMMENTARY_RE.match(line):
                patterns.append(("COMMENTARY", excerpt))
                break
            if _INSTRUCTION_ECHO_RE.search(line):
                patterns.append(("INSTRUCTION_ECHO", excerpt))
                break

    return patterns


def _is_outlier(count: int, median: float) -> bool:
    """Whether a count sits far enough from the median of its unit to be a
    structural problem rather than annotation disagreement."""
    gap = abs(count - median)
    if gap < _MIN_ABSOLUTE_GAP:
        return False
    if median < _SMALL_MEDIAN:
        return gap >= _SMALL_MEDIAN_GAP
    return count > _HIGH_RATIO * median or count < _LOW_RATIO * median


def _detect_count_outliers(records: list[dict]) -> list[dict]:
    """Compare every count against the median across conditions for the same
    text, format and unit."""
    groups: dict[tuple, list[dict]] = defaultdict(list)
    for record in records:
        for unit in UNIT_KEYS:
            count = record["parsed"].get(unit)
            if count is not None:
                groups[(record["text_id"], record["format"], unit)].append(
                    {**record, "count": count})

    outliers = []
    for (text_id, output_format, unit), entries in groups.items():
        if len(entries) < _MIN_CONDITIONS:
            continue
        median = float(np.median([e["count"] for e in entries]))

        for entry in entries:
            count = entry["count"]
            if not _is_outlier(count, median):
                continue
            outliers.append({
                "condition": entry["condition"],
                "format": output_format,
                "text_id": text_id,
                "rep": entry["rep"],
                "attempt": entry["attempt"],
                "pattern": (
                    f"COUNT_OUTLIER_{'HIGH' if count > median else 'LOW'} "
                    f"{unit}:{count} (median {median:.0f})"
                ),
                "parsed": entry["parsed"],
                "output": (_block_excerpt(unit, entry["raw"])
                           if output_format == "lists" else entry["raw"]),
            })

    return outliers


def _finding(phase: str, size: str, corpus_texts: dict[str, str],
             condition: str, **fields) -> dict:
    model, reasoning = _split_condition(condition)
    return {
        "phase": phase,
        "model": model,
        "size": size,
        "reasoning": reasoning,
        "condition": condition,
        "text": corpus_texts.get(fields["text_id"], ""),
        **fields,
    }


def _iter_responses(phase_dir: Path):
    """Yield (condition, format, text_id, rep, attempt, record) for every valid
    saved response under a phase directory."""
    for condition_dir in sorted(d for d in phase_dir.iterdir() if d.is_dir()):
        for format_dir in sorted(d for d in condition_dir.iterdir() if d.is_dir()):
            text_dirs = [d for d in format_dir.iterdir() if d.is_dir()]
            for text_dir in sorted(text_dirs, key=lambda p: int(p.name)):
                for path in sorted(text_dir.glob("rep*_attempt*.json")):
                    with open(path, encoding="utf-8") as f:
                        record = json.load(f)
                    if not record.get("valid"):
                        continue
                    match = re.match(r"rep(\d+)_attempt(\d+)", path.stem)
                    rep, attempt = match.groups() if match else ("?", "?")
                    yield (condition_dir.name, format_dir.name, text_dir.name,
                           rep, attempt, record)


def scan_phase(phase: str, output_dir: Path,
               corpus_texts: dict[str, str]) -> list[dict]:
    """Every non-standard pattern found in one phase."""
    phase_dir = output_dir / phase
    if not phase_dir.exists():
        print(f"  skipping {phase}: {phase_dir} not found", file=sys.stderr)
        return []

    size = _PHASE_SIZES.get(phase, "unknown")
    findings = []
    records = []

    for condition, fmt, text_id, rep, attempt, record in _iter_responses(phase_dir):
        raw = record.get("raw_response", "")
        parsed = parse(raw, fmt)

        if fmt == "counts":
            detected = _detect_counts_patterns(raw)
        elif fmt == "lists":
            detected = _detect_lists_patterns(raw)
        else:
            continue

        for pattern, output in detected:
            findings.append(_finding(
                phase, size, corpus_texts, condition,
                format=fmt, text_id=text_id, rep=rep, attempt=attempt,
                pattern=pattern, parsed=parsed, output=output,
            ))

        if parsed:
            records.append({
                "condition": condition, "format": fmt, "text_id": text_id,
                "rep": rep, "attempt": attempt, "parsed": parsed, "raw": raw,
            })

    findings.extend(
        _finding(phase, size, corpus_texts, **outlier)
        for outlier in _detect_count_outliers(records)
    )
    return findings


def write_xlsx(findings: list[dict], path: Path, overrides: dict) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "nonstandard_responses"

    header_fill = PatternFill(start_color="D9E1F2", fill_type="solid")
    override_fill = PatternFill(start_color="E2EFDA", fill_type="solid")
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")

    for index, name in enumerate(COLUMNS, 1):
        cell = sheet.cell(row=1, column=index, value=name)
        cell.font = Font(bold=True)
        cell.fill = override_fill if name.startswith("override_") else header_fill
        sheet.column_dimensions[get_column_letter(index)].width = (
            _COLUMN_WIDTHS.get(name.removeprefix("override_"), 12))

    for row_index, finding in enumerate(findings, 2):
        parsed = finding.get("parsed") or {}
        corrections = overrides.get((
            finding["phase"], finding["condition"], finding["format"],
            str(finding["text_id"]), int(finding["rep"]),
            int(finding["attempt"]),
        ), {})
        values = [
            finding["phase"], finding["model"], finding["size"],
            finding["reasoning"], finding["format"], finding["condition"],
            int(finding["text_id"]), int(finding["rep"]),
            int(finding["attempt"]), finding["pattern"],
            *[parsed.get(unit) for unit in UNIT_KEYS],
            *[corrections.get(unit) for unit in UNIT_KEYS],
            finding["output"], finding["text"],
        ]
        for column_index, value in enumerate(values, 1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.alignment = wrap if column_index >= len(COLUMNS) - 1 else top

    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Scan model outputs for non-standard response patterns."
    )
    parser.add_argument(
        "--phase", type=str, nargs="+",
        default=[p for p in PHASES if p != "pilot"],
        help="phases to scan (default: every phase but the pilot)",
    )
    parser.add_argument(
        "--output", type=Path, default=None,
        help="output path (default: results/nonstandard_responses.xlsx)",
    )
    parser.add_argument(
        "--overrides", type=Path, default=None,
        help="path to manual_overrides.json, to show the corrected counts "
             "beside the parsed ones",
    )
    args = parser.parse_args()

    out_path = args.output or RESULTS_DIR / "nonstandard_responses.xlsx"
    corpus_texts = _load_texts(CORPUS_PATH)
    overrides = load_overrides(args.overrides) if args.overrides else {}

    findings = []
    for phase in args.phase:
        print(f"Scanning {phase}...", file=sys.stderr)
        phase_findings = scan_phase(phase, OUTPUTS_DIR, corpus_texts)
        findings.extend(phase_findings)
        print(f"  {len(phase_findings)} non-standard patterns found",
              file=sys.stderr)

    write_xlsx(findings, out_path, overrides)
    print(f"\nWrote {len(findings)} rows to {out_path}", file=sys.stderr)


if __name__ == "__main__":
    main()
